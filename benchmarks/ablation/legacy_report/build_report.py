from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "analysis_outputs"
DOCX = ROOT / "VND0_VNDzou_Computational_Study.docx"

NAVY = "17324D"
BLUE = "1D5D8F"
TEAL = "0B7285"
LIGHT_BLUE = "EAF3F8"
PALE_TEAL = "E7F5F3"
LIGHT_GRAY = "F1F3F5"
MID_GRAY = "ADB5BD"
DARK_GRAY = "343A40"
WHITE = "FFFFFF"
RED = "A61E4D"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, value, *, bold=False, color=DARK_GRAY, size=7.4, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_col_widths(table, widths_inches):
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths=None, font_size=7.4, first_col_bold=False, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_text(header.cells[i], h, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header.cells[i], header_fill)
    set_repeat_table_header(header)
    prevent_row_split(header)
    for ridx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.LEFT if cidx in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(row.cells[cidx], value, bold=(first_col_bold and cidx == 0), size=font_size, align=align)
            if ridx % 2:
                set_cell_shading(row.cells[cidx], LIGHT_GRAY)
    if widths:
        set_col_widths(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    return table


def add_caption(doc, label, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Caption"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{label}. ")
    r.bold = True
    p.add_run(text)
    return p


def add_figure(doc, path, width_inches, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    return shape


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_hyperlink(paragraph, text, url, size_half_points=None):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    if size_half_points is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size_half_points))
        r_pr.append(sz)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def fmt(x, digits=2):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "—"
    return f"{float(x):,.{digits}f}"


def fmt_pct(x, digits=2):
    return f"{float(x):.{digits}f}%" if pd.notna(x) else "—"


def fmt_p(p):
    if pd.isna(p):
        return "—"
    p = float(p)
    if p < 0.001:
        return f"{p:.2e}"
    return f"{p:.3f}"


def add_body(doc, text, *, bold_lead=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    if keep:
        p.paragraph_format.keep_together = True
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    # Keep each bullet on one page so a single item never starts with an
    # orphaned line at the foot of the preceding page.
    return add_body(doc, text, style=style, keep=True)


def add_callout(doc, title, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(text)
    return table


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.75)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.85)
    sec.right_margin = Cm(1.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color, before, after in [
        ("Title", 26, NAVY, 0, 8),
        ("Subtitle", 11, TEAL, 0, 8),
        ("Heading 1", 16, NAVY, 14, 6),
        ("Heading 2", 12.3, BLUE, 10, 4),
        ("Heading 3", 10.2, TEAL, 8, 3),
    ]:
        st = doc.styles[style_name]
        st.font.name = "Aptos Display" if "Heading" in style_name or style_name == "Title" else "Aptos"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = style_name != "Subtitle"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    cap = doc.styles["Caption"]
    cap.font.name = "Aptos"
    cap.font.size = Pt(8)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    cap.paragraph_format.space_after = Pt(4)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "VND0–VNDzou computational study  |  deterministic LLRP"
        header.style = doc.styles["Caption"]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        add_page_number(footer)
        footer.style = doc.styles["Caption"]


def add_title_page(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=450, start=350, bottom=450, end=350)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("COMPUTATIONAL STUDY")
    r.font.name = "Aptos"
    r.font.size = Pt(9)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("8FD3C7")
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after = Pt(8)
    r = p2.add_run("VND0 versus VNDzou for the\nLatency Location-Routing Problem")
    r.font.name = "Aptos Display"
    r.font.size = Pt(25)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    r = p3.add_run("Cross-benchmark evaluation, literature positioning, and Barreto ablation study")
    r.font.name = "Aptos"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("D9EAF2")

    doc.add_paragraph()
    add_callout(
        doc,
        "Purpose",
        "A manuscript-ready, evidence-based assessment of solution quality, robustness, runtime, and algorithmic components across the Barreto, Prins/Prodhon, and Tuzun–Burke benchmark sets.",
        fill=PALE_TEAL,
        accent=TEAL,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Analysis date").bold = True
    p.add_run("  22 July 2026\n")
    p.add_run("Study type").bold = True
    p.add_run("  Re-analysis of supplied run-level and aggregate workbooks\n")
    p.add_run("Primary outcome").bold = True
    p.add_run("  Total customer latency (lower is better)")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.add_run("Prepared for direct integration into an academic manuscript. Results are reported with explicit comparability limits and reproducibility notes.").italic = True
    doc.add_page_break()


def main():
    primary = pd.read_csv(OUT / "vnd0_vs_vndzou_summary.csv")
    inst = pd.read_csv(OUT / "vnd0_vs_vndzou_instance.csv")
    inst = inst[inst["_merge"] == "both"].copy()
    run_counts = pd.read_csv(OUT / "run_level_significance_counts.csv")
    lit_comp = pd.read_csv(OUT / "literature_comparison_summary.csv")
    lit_ranks = pd.read_csv(OUT / "literature_average_ranks.csv")
    friedman = pd.read_csv(OUT / "literature_friedman.csv")
    abl = pd.read_csv(OUT / "ablation_summary.csv")
    contrasts = pd.read_csv(OUT / "ablation_component_contrasts.csv")
    size = pd.read_csv(OUT / "vnd0_vs_vndzou_size.csv")

    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    overall_best = primary[(primary.Dataset == "Overall") & (primary.Metric == "best-of-runs")].iloc[0]
    overall_mean = primary[(primary.Dataset == "Overall") & (primary.Metric == "run mean")].iloc[0]

    doc.add_heading("Abstract", level=1)
    abstract = (
        "This study compares a baseline GRASP–VND implementation (VND0) with an extended strategic-oscillation variant (VNDzou) on three deterministic latency location-routing problem benchmark sets. "
        "The paired primary analysis covers 75 instances: 10 Barreto, 30 Prins/Prodhon, and 35 Tuzun–Burke. VNDzou improved run-mean latency on all 75 matched instances, with an overall mean improvement of "
        f"{overall_mean['Mean improvement (%)']:.2f}% (95% bootstrap CI {overall_mean['Mean improvement 95% CI lower']:.2f}%–{overall_mean['Mean improvement 95% CI upper']:.2f}%; Holm-adjusted Wilcoxon p={fmt_p(overall_mean['Wilcoxon p (Holm)'])}). "
        f"For best-of-runs performance, VNDzou recorded {int(overall_best['VNDzou wins'])} wins and {int(overall_best['Ties'])} ties with no losses, while increasing median CPU time by {overall_best['Median CPU ratio Z/0']:.2f}×. "
        "Against published deterministic results, VNDzou strongly dominated RGA, MA, and GBILS and was competitive with SA–VND variants, but remained inferior to M-ILS and especially RLHEA on average performance. "
        "A 13-configuration Barreto study supports a mean-quality contribution from strategic oscillation and shows that the full VNDzou configuration has the strongest complete-design quality, while runtime trade-offs, incomplete configurations, and naming-to-code inconsistencies constrain broader causal interpretation."
    )
    add_body(doc, abstract)
    p = doc.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run("latency location-routing problem; variable neighborhood descent; strategic oscillation; algorithm comparison; ablation study; reproducibility")

    doc.add_heading("Key findings", level=1)
    key_rows = [
        ["Quality", "75/75", "VNDzou wins on run-mean latency across every matched instance."],
        ["Best-of-runs", "61/14/0", "VNDzou/VND0 win–tie–loss record at two-decimal precision."],
        ["Runtime", f"{overall_best['Median CPU ratio Z/0']:.2f}×", "Median CPU multiplier of VNDzou relative to VND0."],
        ["Literature", "RLHEA leads", "VNDzou is not a reproduction of RLHEA and does not reach its average quality."],
        ["Ablation", "SO matters", "Strategic oscillation improves all ten Barreto run means; the full VNDzou configuration yields the strongest overall quality."],
    ]
    add_table(doc, ["Dimension", "Result", "Interpretation"], key_rows, widths=[1.15, 1.05, 4.55], font_size=8.2, first_col_bold=True, header_fill=TEAL)
    add_callout(
        doc,
        "Manuscript claim supported by the evidence",
        "VNDzou is a materially stronger and more stable intensification mechanism than VND0 across the tested deterministic LLRP instances, at an approximate threefold median runtime cost. It is competitive with—but does not supersede—the strongest published hybrid evolutionary method.",
    )

    doc.add_heading("1. Study scope and research questions", level=1)
    add_body(doc, "The analysis addresses four research questions:")
    add_bullet(doc, "RQ1 — Does VNDzou improve solution quality and robustness relative to VND0 across Barreto, Prins/Prodhon, and Tuzun–Burke instances?")
    add_bullet(doc, "RQ2 — What computational-time premium accompanies those quality gains?")
    add_bullet(doc, "RQ3 — Where does VNDzou sit relative to directly comparable deterministic LLRP methods from the literature?")
    add_bullet(doc, "RQ4 — Which Barreto configuration changes are associated with quality or runtime differences, and which conclusions remain limited by design or data quality?")

    doc.add_heading("1.1 Problem and outcome definition", level=2)
    add_body(doc, "The evaluated code minimizes total customer latency: the sum of customer arrival times over open depot-to-customer routes. Lower values are better. The supplied implementations do not count a return-to-depot leg in the objective. Feasibility requires each customer to be visited once, route capacity compliance, the prescribed fleet limit, and the required number of open depots.")

    doc.add_heading("1.2 Method identity and comparability boundary", level=2)
    add_body(doc, "VND0 and VNDzou share the same randomized ECWA construction, GRASP framework (50 iterations; five diversification iterations followed by three intensification iterations), and feasibility model. Their principal difference is the local-search engine.")
    algo_rows = [
        ["Construction / outer search", "ECWA + GRASP", "ECWA + GRASP", "Controlled"],
        ["Route neighborhoods", "5: insertion, swap, 2-opt, arc-swap, shift2-1", "7: relocate, swap, 2-opt, 2-relocate, node–arc swap, arc–arc swap, SwapStar", "Expanded in VNDzou"],
        ["Capacity handling", "Feasible-only local moves", "Adaptive penalty with strategic oscillation", "Expanded in VNDzou"],
        ["Depot search", "Construction-driven", "Depot-swap perturbation", "Expanded in VNDzou"],
        ["Population / crossover", "None", "None", "Unlike RLHEA"],
        ["Q-learning in base method", "None", "None", "Unlike RLHEA"],
    ]
    add_caption(doc, "Table 1", "Algorithmic scope of the two implementations.")
    add_table(doc, ["Component", "VND0", "VNDzou", "Interpretive status"], algo_rows, widths=[1.25, 1.9, 2.25, 1.35], font_size=7.5, first_col_bold=True)
    add_callout(doc, "Terminology safeguard", "The name “VNDzou” denotes the supplied Zou-inspired GRASP–VND–strategic-oscillation implementation. It is not an implementation of the population-based RLHEA of Zou, Hao, and Wu (2024), which also uses multi-parent edge-assembly crossover and reinforcement learning.", fill="FFF4E6", accent="C05621")

    doc.add_heading("2. Experimental design and statistical protocol", level=1)
    coverage_rows = [
        ["Barreto", "10", "10", "10", "30 / 30", "Complete"],
        ["Prins/Prodhon", "30", "30", "30", "50 / 30*", "29 VND0 instances have 50 runs; one has 30"],
        ["Tuzun–Burke", "36", "35", "35", "30 / 30", "P132222-20 lacks VND0 results"],
        ["Total", "76", "75", "75", "Unequal in Prins", "Primary paired n=75"],
    ]
    add_caption(doc, "Table 2", "Coverage and stochastic-replication budgets.")
    add_table(doc, ["Set", "VNDzou", "VND0", "Matched", "Runs VND0 / VNDzou", "Note"], coverage_rows, widths=[1.2, 0.65, 0.65, 0.65, 1.35, 2.25], font_size=7.4, first_col_bold=True)
    add_body(doc, "* VND0 has 50 completed runs for 29 Prins instances; 200-10-3b-10 has 30. Consequently, run means are the primary distributional comparison, while best-of-runs results are reported with an explicit budget caveat.")

    doc.add_heading("2.1 Estimands and tests", level=2)
    add_body(doc, "For instance i, the reported relative improvement is 100 × (VND0ᵢ − VNDzouᵢ) / VND0ᵢ; positive values favor VNDzou. Win–tie–loss counts use objective values rounded to two decimals, matching the precision of published benchmark tables.")
    add_bullet(doc, "Primary inference: two-sided paired Wilcoxon signed-rank tests across instances, with Holm correction within each outcome family; matched-pairs rank-biserial correlation quantifies effect direction and magnitude.")
    add_bullet(doc, "Uncertainty: nonparametric percentile bootstrap confidence intervals for the mean instance-level relative improvement (20,000 resamples; fixed seed).")
    add_bullet(doc, "Supportive run-level analysis: two-sided Mann–Whitney tests per instance, Holm-adjusted within benchmark set, with Cliff’s delta. The instance remains the unit of generalization.")
    add_bullet(doc, "Robustness: coefficient of variation (CV = standard deviation / mean) and its VNDzou/VND0 ratio.")
    add_bullet(doc, "Reference quality: best-known-solution (BKS) matches are evaluated at two-decimal precision against the RLHEA benchmark table.")

    doc.add_heading("2.2 Runtime interpretation", level=2)
    add_body(doc, "CPU comparisons between VND0 and VNDzou are descriptive within the supplied experimental suite. Processor, Python, and node-allocation metadata were not preserved in the result folder, so a shared environment is assumed but cannot be independently audited. Published CPU times are not subjected to inferential comparison because languages, hardware, stopping criteria, and run budgets differ.")

    doc.add_heading("3. VND0 versus VNDzou", level=1)
    doc.add_heading("3.1 Solution quality", level=2)
    rows = []
    for ds in ["Barreto", "Prins", "Tuzun", "Overall"]:
        for metric in ["best-of-runs", "run mean"]:
            r = primary[(primary.Dataset == ds) & (primary.Metric == metric)].iloc[0]
            rows.append([
                ds,
                "Best" if metric == "best-of-runs" else "Mean",
                int(r["n"]),
                f"{int(r['VNDzou wins'])}/{int(r['Ties'])}/{int(r['VND0 wins'])}",
                fmt_pct(r["Mean improvement (%)"]),
                f"[{r['Mean improvement 95% CI lower']:.2f}, {r['Mean improvement 95% CI upper']:.2f}]",
                fmt_p(r["Wilcoxon p (Holm)"]),
                f"{r['Rank-biserial (positive favors VNDzou)']:.3f}",
            ])
    add_caption(doc, "Table 3", "Paired solution-quality comparison. W/T/L is VNDzou versus VND0.")
    add_table(doc, ["Set", "Outcome", "n", "W/T/L", "Mean Δ", "95% CI", "pHolm", "rᵣᵦ"], rows, widths=[1.0, 0.7, 0.4, 0.75, 0.75, 1.15, 0.8, 0.6], font_size=7.2, first_col_bold=True)

    add_body(doc, f"Across all 75 matched instances, VNDzou reduced run-mean latency by {overall_mean['Mean improvement (%)']:.2f}% on average (95% CI {overall_mean['Mean improvement 95% CI lower']:.2f}%–{overall_mean['Mean improvement 95% CI upper']:.2f}%). Every instance favored VNDzou, and the rank-biserial effect was maximal (rᵣᵦ=1.00). Best-of-runs latency improved by {overall_best['Mean improvement (%)']:.2f}% on average, with 61 wins, 14 ties, and no VND0 wins.")
    add_body(doc, "The gain increases with instance scale: mean run-mean improvements are 0.97% for ≤50-customer instances, 6.11% for 51–100 customers, and 8.07% for >100 customers. The strongest gains occur in the Tuzun–Burke set, where several high-variance VND0 outcomes are substantially stabilized by VNDzou.")
    add_caption(doc, "Figure 1", "Instance-level change in run-mean latency; positive values favor VNDzou.")
    add_figure(doc, OUT / "fig_primary_improvement.png", 6.65, "Three horizontal bar charts show VNDzou's percent improvement in mean latency over VND0 for Barreto, Prins, and Tuzun instances. Most bars are positive and Tuzun has the largest gains.")

    doc.add_heading("3.2 Runtime and stability", level=2)
    cpu = inst.groupby("Dataset").agg(
        cpu0=("CPU_0", "mean"), cpuz=("CPU_Z", "mean"),
        medratio=("CPU ratio Z/0", "median"), meanratio=("CPU ratio Z/0", "mean"),
        cv0=("CV_0", "median"), cvz=("CV_Z", "median"), cvratio=("CV ratio Z/0", "median"),
    ).reset_index()
    cpu_rows = []
    for _, r in cpu.iterrows():
        cpu_rows.append([r.Dataset, fmt(r.cpu0, 1), fmt(r.cpuz, 1), f"{r.medratio:.2f}×", f"{r.meanratio:.2f}×", f"{r.cv0:.4f}", f"{r.cvz:.4f}", f"{r.cvratio:.3f}"])
    add_caption(doc, "Table 4", "Within-study runtime and run-to-run variability.")
    add_table(doc, ["Set", "CPU₀", "CPUZ", "Median ratio", "Mean ratio", "Median CV₀", "Median CVZ", "CV ratio"], cpu_rows, widths=[1.0, 0.8, 0.8, 0.95, 0.85, 0.9, 0.9, 0.7], font_size=7.2, first_col_bold=True)
    add_body(doc, f"The median CPU multiplier is {overall_best['Median CPU ratio Z/0']:.2f}× overall, increasing from 2.42× on Barreto to 3.45× on Tuzun–Burke. In return, VNDzou's median CV is less than half that of VND0 overall (median CV ratio {overall_best['Median CV ratio Z/0']:.3f}), indicating materially tighter run-to-run outcomes.")
    add_caption(doc, "Figure 2", "Benchmark-level quality–runtime trade-off.")
    add_figure(doc, OUT / "fig_quality_runtime.png", 6.2, "Scatter plot positions Barreto, Prins, and Tuzun by VNDzou-to-VND0 CPU ratio and mean latency improvement. Tuzun has the largest quality gain and runtime multiplier.")

    doc.add_heading("3.3 Run-level distributional evidence", level=2)
    run_rows = []
    for _, r in run_counts.iterrows():
        run_rows.append([r.Dataset, int(r.get("VNDzou", 0)), int(r.get("VND0", 0)), int(r.get("Not significant", 0)), int(r.drop(labels="Dataset").sum())])
    add_caption(doc, "Table 5", "Direction of Holm-adjusted per-instance Mann–Whitney tests.")
    add_table(doc, ["Set", "VNDzou significantly better", "VND0 significantly better", "Not significant", "Tested"], run_rows, widths=[1.2, 1.65, 1.55, 1.1, 0.7], font_size=7.5, first_col_bold=True)
    add_body(doc, "VNDzou is significantly better in 68 of 73 instances with available raw results (7/10 Barreto, 28/30 Prins, and 33/33 Tuzun–Burke); no instance significantly favors VND0. This supports the paired instance-level result without treating hundreds of stochastic runs as independent benchmark replications.")

    doc.add_heading("3.4 Best-known-solution attainment", level=2)
    bks_rows = []
    for ds, g in inst.groupby("Dataset"):
        bks_rows.append([ds, len(g), int(g["BKS hit VND0"].sum()), int(g["BKS hit VNDzou"].sum()), fmt_pct(g["Gap BKS VND0 (%)"].mean()), fmt_pct(g["Gap BKS VNDzou (%)"].mean())])
    bks_rows.append(["Overall", len(inst), int(inst["BKS hit VND0"].sum()), int(inst["BKS hit VNDzou"].sum()), fmt_pct(inst["Gap BKS VND0 (%)"].mean()), fmt_pct(inst["Gap BKS VNDzou (%)"].mean())])
    add_caption(doc, "Table 6", "Attainment of the published RLHEA benchmark value.")
    add_table(doc, ["Set", "n", "VND0 hits", "VNDzou hits", "Mean best gap VND0", "Mean best gap VNDzou"], bks_rows, widths=[1.1, 0.5, 0.8, 0.85, 1.35, 1.45], font_size=7.5, first_col_bold=True)
    add_body(doc, "VNDzou raises the number of two-decimal BKS matches from 17 to 24 among the 75 matched instances. These counts describe benchmark attainment, not proven optimality: most Tuzun–Burke values remain best-known upper bounds.")

    doc.add_heading("4. Position relative to the literature", level=1)
    doc.add_heading("4.1 Directly comparable methods", level=2)
    add_body(doc, "The deterministic benchmark tables permit objective-value comparison with RGA and MA (Moshref-Javadi and Lee, 2016), GBILS (Nucamendi-Guillén et al., 2022), SA–VND0/1/2 (Osorio-Mora et al., 2023b), M-ILS (Osorio-Mora et al., 2023a), and RLHEA (Zou et al., 2024). Deterministic MS-VNS results reported in the supplement of Osorio-Mora et al. (2025) are included as a secondary five-run comparator.")
    add_body(doc, "The 2026 split-delivery/opening-cost variant is excluded from numerical ranking because its objective and feasibility structure differ. The same scope restriction applies to the stochastic travel-time model in the main 2025 paper; only its deterministic supplementary tables are used.")

    overall_lit = lit_comp[lit_comp.Dataset == "Overall"].copy()
    order = ["RGA", "MA", "GBILS", "SA-VND0", "SA-VND1", "SA-VND2", "M-ILS", "RLHEA", "MS-VNS"]
    lit_rows = []
    for method in order:
        best = overall_lit[(overall_lit.Method == method) & (overall_lit.Metric == "best-of-runs")]
        mean = overall_lit[(overall_lit.Method == method) & (overall_lit.Metric == "run mean")]
        b = best.iloc[0] if len(best) else None
        m = mean.iloc[0] if len(mean) else None
        lit_rows.append([
            method,
            int(b.n) if b is not None else "—",
            f"{int(b['VNDzou wins'])}/{int(b['Ties'])}/{int(b['Literature wins'])}" if b is not None else "—",
            fmt_pct(b["Mean VNDzou improvement (%)"]) if b is not None else "—",
            fmt_p(b["Wilcoxon p (Holm)"]) if b is not None else "—",
            f"{int(m['VNDzou wins'])}/{int(m['Ties'])}/{int(m['Literature wins'])}" if m is not None else "—",
            fmt_pct(m["Mean VNDzou improvement (%)"]) if m is not None else "—",
            fmt_p(m["Wilcoxon p (Holm)"]) if m is not None else "—",
        ])
    add_caption(doc, "Table 7", "Overall literature comparison. W/T/L is VNDzou versus the named method.")
    add_table(doc, ["Method", "n", "Best W/T/L", "Best Δ", "pHolm", "Mean W/T/L", "Mean Δ", "pHolm"], lit_rows, widths=[1.05, 0.4, 0.85, 0.7, 0.7, 0.85, 0.7, 0.7], font_size=6.9, first_col_bold=True)

    doc.add_heading("4.2 Interpretation", level=2)
    add_body(doc, "VNDzou decisively improves on the first-generation RGA and MA baselines: it wins all 76 run-mean comparisons and improves the mean objective by 12.85% and 10.09%, respectively. It also records 29 wins and nine ties against GBILS best values over the 38 instances for which GBILS results are available.")
    add_body(doc, "Against SA–VND0/1/2, VNDzou has a statistically stronger best-of-runs profile overall, but average solution quality is effectively comparable: none of the three overall run-mean contrasts remains significant after Holm correction. This distinction is important—VNDzou reaches strong solutions more often than the SA–VND variants without producing a uniformly better 30-run mean.")
    add_body(doc, "M-ILS and RLHEA remain stronger. Relative to RLHEA, VNDzou records 9 wins, 15 ties, and 52 losses in best-of-runs results, and 2 wins, 5 ties, and 69 losses in run means. On average, VNDzou is 0.61% worse for the best value and 2.20% worse for the run mean. The RLHEA gap is significant after Holm correction (best p=5.03×10⁻⁷; mean p=3.06×10⁻¹²).")
    add_body(doc, "The deterministic MS-VNS comparison is mixed and not statistically decisive overall: best values split 35/6/35 and run means 34/2/40. Its five-run budget and different hardware make this a secondary positioning result.")
    add_caption(doc, "Figure 3", "Available-case average ranks for run-mean solution quality.")
    add_figure(doc, OUT / "fig_literature_ranks.png", 6.15, "Horizontal bars show average ranks of literature methods for mean solution quality. RLHEA ranks first, M-ILS second, and VNDzou is in the middle of the modern methods.")
    fr_best = friedman[friedman.Metric == "best-of-runs"].iloc[0]
    fr_mean = friedman[friedman.Metric == "run mean"].iloc[0]
    add_body(doc, f"A complete-case Friedman test across VNDzou, the three SA–VND methods, M-ILS, and RLHEA rejects equal performance for both best values (χ²={fr_best['Friedman chi-square']:.2f}, p={fmt_p(fr_best.p)}) and run means (χ²={fr_mean['Friedman chi-square']:.2f}, p={fmt_p(fr_mean.p)}). Among these six methods, VNDzou's average rank is 3.36 for best-of-runs and 4.21 for run mean; RLHEA ranks 1.68 and 1.28, respectively.")

    doc.add_heading("4.3 Why CPU times are not cross-compared", level=2)
    hw_rows = [
        ["RGA / MA", "30", "3.1 GHz computer; 4 GB RAM", "Original study"],
        ["GBILS", "5", "Intel Core i5-6300, 2.40 GHz; 8 GB; C++", "Original study"],
        ["SA–VND / M-ILS / RLHEA", "30", "Xeon E5-2670, 2.5 GHz; 2 GB; Linux; single thread; C++ -O3 for RLHEA", "Re-run framework in Zou et al. (2024)"],
        ["MS-VNS", "5", "AMD Ryzen 7 2700X, 3.7 GHz; 64 GB; Linux; single thread; C++", "2025 supplement"],
        ["VND0 / VNDzou", "30–50", "Execution metadata absent; Python source supplied", "Current result archive"],
    ]
    add_caption(doc, "Table 8", "Reported experimental environments and run budgets.")
    add_table(doc, ["Method", "Runs", "Environment", "Source"], hw_rows, widths=[1.25, 0.55, 3.35, 1.55], font_size=7.2, first_col_bold=True)

    doc.add_heading("5. Barreto ablation study", level=1)
    doc.add_heading("5.1 Design and limits", level=2)
    add_body(doc, "Thirteen supplied configuration summaries were evaluated. Twelve cover all ten Barreto instances; VNDZou-CAND30 covers only seven and is excluded from complete-design ranks and Pareto labeling. The archive is not a balanced full factorial: component effects are interpreted only where two configurations differ by one named change, and even then code provenance and multiple testing limit causal claims.")

    abl_sorted = abl.sort_values(["n", "Mean best gap (%)"], ascending=[False, True])
    abl_rows = []
    for _, r in abl_sorted.iterrows():
        abl_rows.append([
            r.Configuration,
            int(r.n),
            fmt_pct(r["Mean best gap (%)"]),
            fmt_pct(r["Mean run-mean gap (%)"]),
            fmt(r["Median CPU"], 1),
            int(r["BKS hits"]),
            fmt(r["Complete-design best rank"], 2),
            fmt(r["Complete-design mean rank"], 2),
            "Yes" if bool(r["Pareto efficient"]) else "No",
        ])
    add_caption(doc, "Table 9", "Barreto configuration performance. Complete-design ranks exclude VNDZou-CAND30.")
    add_table(doc, ["Configuration", "n", "Best gap", "Mean gap", "Median CPU", "BKS", "Best rank", "Mean rank", "Pareto"], abl_rows, widths=[1.7, 0.35, 0.65, 0.7, 0.75, 0.4, 0.65, 0.7, 0.55], font_size=6.5, first_col_bold=True)
    add_body(doc, "The base VNDzou configuration is the strongest complete configuration on average: mean best gap 0.46%, mean run-mean gap 1.85%, and the best complete-design run-mean rank (1.85). Its median CPU time is 86.6 s. VND0-SO-DEPOT offers a lower-cost compromise (0.79% best gap; 2.21% mean gap; 53.6 s median CPU), while the H3 construction is very fast but loses substantial quality.")
    add_caption(doc, "Figure 4", "Quality–runtime frontier for complete Barreto configurations.")
    add_figure(doc, OUT / "fig_ablation_pareto.png", 6.35, "Log-scaled scatter plot compares median CPU time and average best-known-solution gap for complete Barreto ablation configurations. Teal points are non-dominated.")

    doc.add_heading("5.2 Component contrasts", level=2)
    contrast_rows = []
    for _, r in contrasts.iterrows():
        contrast_rows.append([
            r.Contrast,
            int(r.n),
            fmt_pct(r["Best improvement (%)"]),
            str(r["Best W/T/L"]),
            fmt_p(r["Best p (Holm)"]),
            fmt_pct(r["Mean improvement (%)"]),
            str(r["Mean W/T/L"]),
            fmt_p(r["Mean p (Holm)"]),
            f"{r['Median CPU ratio variant/base']:.2f}×",
        ])
    add_caption(doc, "Table 10", "Matched component contrasts; positive improvement favors the variant.")
    add_table(doc, ["Contrast", "n", "Best Δ", "W/T/L", "pHolm", "Mean Δ", "W/T/L", "pHolm", "CPU ratio"], contrast_rows, widths=[2.05, 0.3, 0.55, 0.55, 0.6, 0.55, 0.55, 0.6, 0.7], font_size=6.2, first_col_bold=True)
    add_bullet(doc, "Strategic oscillation is the cleanest positive VND0 contrast: VND0-SO improves the run mean on all ten instances by 0.40% on average (Holm p=0.018) at a median 1.15× CPU cost; best values do not materially change.")
    add_bullet(doc, "Adding depot-swap to VND0-SO improves eight of ten run means by 0.72% with almost unchanged median CPU (1.00×); the Holm-adjusted p-value is 0.068, so the evidence is directional rather than conventionally significant.")
    add_bullet(doc, "The VND0 CAND20 variant reduces median CPU to 0.83× the baseline without a systematic quality change. Adaptive or Q-guided ordering adds no robust gain on this ten-instance design.")
    add_bullet(doc, "The VNDzou-SO-DEPOT-CAND20 folder improves nine run means and ties one relative to VNDZou-SO-DEPOT (1.28% mean gain; Holm p=0.031), but requires 4.01× median CPU. The checked operator source sets CANDIDATE_DELTA=30, so the folder label “CAND20” is not an auditable parameter claim.")
    add_bullet(doc, "The H3 construction cuts median CPU to 0.24× VND0 but worsens the mean objective by 9.37%; after Holm correction the small-n test is not significant, although the descriptive loss is large and consistent in eight of ten best values.")

    doc.add_heading("6. Discussion", level=1)
    doc.add_heading("6.1 What the evidence supports", level=2)
    add_body(doc, "The central result is unusually consistent: VNDzou improves every matched run mean and never loses a best-of-runs comparison to VND0. The gain is small on the easiest/smallest instances, where both methods frequently reproduce the same two-decimal best value, and grows sharply on larger Tuzun–Burke cases. Simultaneously lower CV values indicate that the improvement is not driven solely by a few exceptional seeds; VNDzou also narrows the outcome distribution.")
    add_body(doc, "The algorithmic reading is that expanded inter-route neighborhoods, controlled traversal of infeasible capacity states, and depot-level perturbation jointly strengthen intensification. The Barreto ablation isolates a modest but reliable mean-quality contribution from strategic oscillation. It does not, however, identify a unique contribution for every VNDzou component because the supplied configuration grid is incomplete and some labels do not match checked code constants.")

    doc.add_heading("6.2 What the evidence does not support", level=2)
    add_bullet(doc, "It does not support calling VNDzou state of the art: RLHEA is significantly stronger, particularly in average quality.")
    add_bullet(doc, "It does not support equating VNDzou with Zou et al.'s RLHEA: the base implementation has no population, MPEAX crossover, or Q-learning neighborhood ordering.")
    add_bullet(doc, "It does not establish hardware-normalized speed relative to literature methods or absolute scalability beyond 200 customers.")
    add_bullet(doc, "It does not fully separate the seven-neighborhood, strategic-oscillation, and depot-perturbation effects; a balanced factorial design is still required.")

    doc.add_heading("6.3 Threats to validity", level=2)
    threats = [
        ["Unequal run budget", "VND0 uses 50 runs for 29 Prins instances versus 30 for VNDzou.", "Prioritize run means; report best values with caveat."],
        ["Missing VND0 instance", "P132222-20 is absent.", "Primary paired analysis uses n=75; literature analysis uses all 76 VNDzou instances."],
        ["Incomplete raw data", "Two Tuzun VND0 instances are summary-only.", "Run-level tests use 33 Tuzun instances; paired summary uses 35."],
        ["Execution metadata", "CPU hardware/interpreter metadata are absent.", "Limit runtime inference to descriptive within-suite ratios."],
        ["Configuration naming", "CAND20 folder source declares δ=30; VNDzou outputs may carry VND0 labels.", "Use file provenance, disclose mismatch, avoid parameter-specific causal claims."],
        ["Multiple comparisons", "Ablation contains nine planned contrasts at n≤10.", "Apply Holm correction and emphasize effect sizes/directions."],
    ]
    add_caption(doc, "Table 11", "Validity threats and mitigations.")
    add_table(doc, ["Threat", "Observed issue", "Mitigation"], threats, widths=[1.25, 2.65, 2.8], font_size=7.1, first_col_bold=True)

    doc.add_heading("7. Conclusions and manuscript-ready statements", level=1)
    add_body(doc, "VNDzou provides a clear quality and robustness improvement over VND0 on the supplied deterministic LLRP benchmark archive. Across 75 matched instances, it improves mean latency in every case and best latency in 61 cases while tying the remaining 14. The overall mean improvements are 5.68% for run means and 1.69% for best-of-runs values; both are statistically compelling and accompanied by lower run-to-run variability. The cost is a median 2.89× increase in CPU time.")
    add_body(doc, "The literature position is credible but bounded. VNDzou dominates RGA, MA, and GBILS and improves the best-of-runs profile relative to SA–VND0/1/2. Its average quality is broadly comparable to the SA–VND family, but M-ILS and RLHEA remain stronger, with RLHEA leading decisively. Accordingly, VNDzou should be presented as an effective, robust GRASP–VND intensification variant rather than as an RLHEA reproduction or a new state-of-the-art solver.")
    add_body(doc, "The Barreto ablation supports strategic oscillation as a useful component and identifies VND0-SO-DEPOT as a favorable lower-cost compromise. The strongest complete quality comes from the base VNDzou configuration. A publication-grade follow-up should run a balanced factorial experiment with identical 30-run budgets, pinned code commits, machine metadata, and verified parameter manifests.")
    add_callout(doc, "Recommended one-sentence result", "Across 75 matched deterministic LLRP instances, VNDzou reduced mean latency by 5.68% (95% bootstrap CI 4.17%–7.46%) and improved or tied every best-of-runs value, while requiring a median 2.89-fold increase in CPU time.", fill=PALE_TEAL, accent=TEAL)

    doc.add_heading("References", level=1)
    references = [
        ("Barreto, S. (2004). Análise e modelização de problemas de localização-distribuição [Analysis and modelling of location-routing problems] (Doctoral dissertation). University of Aveiro.", None),
        ("Cavagnini, R., Santini, A., & Schneider, M. (2026). Recent developments in location-routing problems: Deterministic single-echelon, single-objective, single-period problems. European Journal of Operational Research. ", "https://doi.org/10.1016/j.ejor.2025.09.040"),
        ("Moshref-Javadi, M., & Lee, S. (2016). The latency location-routing problem. European Journal of Operational Research, 255(2), 604–619. ", "https://doi.org/10.1016/j.ejor.2016.05.048"),
        ("Nucamendi-Guillén, S., Martínez-Salazar, I., Khodaparasti, S., & Bruni, M. E. (2022). New formulations and solution approaches for the latency location routing problem. Computers & Operations Research, 143, 105767. ", "https://doi.org/10.1016/j.cor.2022.105767"),
        ("Nucamendi-Guillén, S., Gómez-Rocha, J. E., Camacho-Vallejo, J.-F., Moraga Pardo, J. A., & González-Ramírez, R. G. (2026). The latency location routing problem with split deliveries: Mathematical formulation and metaheuristic algorithm. International Transactions in Operational Research, 33, 1609–1635. ", "https://doi.org/10.1111/itor.13606"),
        ("Osorio-Mora, A., Escobar, J. W., & Toth, P. (2023a). An iterated local search algorithm for latency vehicle routing problems with multiple depots. Computers & Operations Research, 158, 106293.", None),
        ("Osorio-Mora, A., Rey, C., Toth, P., & Vigo, D. (2023b). Effective metaheuristics for the latency location routing problem. International Transactions in Operational Research, 30, 3801–3832. ", "https://doi.org/10.1111/itor.13294"),
        ("Osorio-Mora, A., Saldanha-da-Gama, F., & Toth, P. (2025). A risk-averse latency location-routing problem with stochastic travel times. European Journal of Operational Research, 321, 837–850. ", "https://doi.org/10.1016/j.ejor.2024.10.041"),
        ("Prins, C., Prodhon, C., & Wølfler Calvo, R. (2004). Nouveaux algorithmes pour le problème de localisation et routage avec contraintes de capacité. MOSIM'04, Nantes, France.", None),
        ("Tuzun, D., & Burke, L. I. (1999). A two-phase tabu search approach to the location routing problem. European Journal of Operational Research, 116(1), 87–99.", None),
        ("Zou, Y., Hao, J.-K., & Wu, Q. (2024). A reinforcement learning guided hybrid evolutionary algorithm for the latency location routing problem. Computers & Operations Research, 170, 106758. ", "https://doi.org/10.1016/j.cor.2024.106758"),
    ]
    for text, url in references:
        p = doc.add_paragraph(style="Bibliography" if "Bibliography" in [s.name for s in doc.styles] else None)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 0.95
        rr = p.add_run(text)
        rr.font.name = "Aptos"
        rr.font.size = Pt(8.1)
        if url:
            add_hyperlink(p, url.replace("https://doi.org/", "doi:"), url, size_half_points=16)

    # Landscape appendix with the full matched-instance table.
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.top_margin = Cm(1.35)
    sec.bottom_margin = Cm(1.25)
    sec.left_margin = Cm(1.35)
    sec.right_margin = Cm(1.35)
    doc.add_heading("Appendix A. Full VND0–VNDzou instance comparison", level=1)
    add_body(doc, "Positive Δ values favor VNDzou. BKS is the RLHEA best value reported in the consolidated literature table; “hit” uses two-decimal precision.")
    full_rows = []
    for _, r in inst.sort_values(["Dataset", "Customers", "Instance"]).iterrows():
        full_rows.append([
            r.Dataset,
            r.Instance,
            fmt(r["Benchmark BKS"], 2),
            fmt(r.Best_0, 2),
            fmt(r.Best_Z, 2),
            fmt_pct(r["Best improvement (%)"], 2),
            fmt(r.Mean_0, 2),
            fmt(r.Mean_Z, 2),
            fmt_pct(r["Mean improvement (%)"], 2),
            f"{r['CPU ratio Z/0']:.2f}×",
            fmt_pct(r["Gap BKS VNDzou (%)"], 2),
        ])
    add_table(doc, ["Set", "Instance", "BKS", "Best₀", "BestZ", "Δ best", "Mean₀", "MeanZ", "Δ mean", "CPU Z/0", "Z gap"], full_rows,
              widths=[0.8, 1.25, 0.8, 0.8, 0.8, 0.65, 0.85, 0.85, 0.7, 0.7, 0.65], font_size=6.1, first_col_bold=True)

    # Return to portrait.
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(1.75)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.85)
    sec.right_margin = Cm(1.85)
    doc.add_heading("Appendix B. Reproducibility and data-integrity audit", level=1)
    audit_rows = [
        ["Primary coverage", "VNDzou 76; VND0 75; matched 75", "P132222-20 absent for VND0."],
        ["Raw-run coverage", "Barreto 10; Prins 30; Tuzun 33", "P132212-10 and P133222-20 use aggregate VND0 summaries only."],
        ["Prins run budget", "VND0: 29×50 + 1×30; VNDzou: 30×30", "Best-of-runs is not budget-balanced."],
        ["Feasibility", "100% in all supplied aggregate summaries", "No infeasible terminal run affects reported means."],
        ["Algorithm label", "VNDzou logged output may state GRASP-VND0-SO-LLRP", "Folder/source provenance was used for method identity."],
        ["Candidate delta", "VNDZou-SO-DEPOT-CAND20 operator file sets CANDIDATE_DELTA=30", "Treat folder name as a label, not verified δ=20."],
        ["Incomplete ablation", "VNDZou-CAND30 has 7/10 Barreto instances", "Excluded from complete-design ranks and Pareto frontier."],
        ["Hardware metadata", "Absent for current experiments", "No cross-literature CPU inference."],
        ["Code identity", "Base VNDzou lacks population, MPEAX, and Q-learning", "Do not describe it as RLHEA replication."],
    ]
    add_table(doc, ["Audit item", "Finding", "Consequence"], audit_rows, widths=[1.35, 2.7, 2.55], font_size=7.2, first_col_bold=True)

    doc.add_heading("Appendix C. Reporting checklist for a follow-up experiment", level=1)
    checklist = [
        "Use the same 30 seeds for every algorithm and configuration; preserve a seed manifest.",
        "Pin a Git commit hash and generate configuration names directly from a machine-readable parameter file.",
        "Record CPU model, node, operating system, Python/compiler version, process/thread count, and wall-clock/CPU semantics.",
        "Complete the missing P132222-20 VND0 run and the three incomplete VNDZou-CAND30 instances.",
        "Run a balanced factorial design for neighborhood set × strategic oscillation × depot perturbation × candidate delta × operator ordering.",
        "Pre-specify the primary endpoint as run-mean latency, with best-of-runs secondary; report paired bootstrap intervals and multiplicity control.",
        "Validate every generated best solution independently and archive route/depot decisions alongside objective values.",
    ]
    for item in checklist:
        add_bullet(doc, item)

    doc.core_properties.title = "VND0 versus VNDzou for the Latency Location-Routing Problem"
    doc.core_properties.subject = "Cross-benchmark computational evaluation, literature comparison, and Barreto ablation study"
    doc.core_properties.author = "Computational Research Report"
    doc.core_properties.keywords = "LLRP, VND, GRASP, strategic oscillation, ablation, RLHEA"
    doc.core_properties.comments = "Generated from the supplied computational archive; see Appendix B for data-integrity notes."
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
